"""Build the macro-free CRM_CONTACTOS_v1.xlsx template with openpyxl.

This is an authoring tool only. The API serves the checked-in XLSX binary.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path(__file__).resolve().parents[1] / "apps/api/data/db/CRM_CONTACTOS_v1.xlsx"
HEADERS = (
    "DNI", "Nombre", "Apellido", "Teléfono", "Estado", "Actividad",
    "Fecha de inscripción", "Fecha de vencimiento",
)
STATES = ("Al día", "Adeudando", "Nuevo Inscripto", "Abandonado")


def build() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CONTACTOS_CRM"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H10001"
    sheet.sheet_view.showGridLines = False

    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    for column, width in {"A": 16, "B": 23, "C": 23, "D": 22, "E": 22,
                          "F": 27, "G": 25, "H": 25, "J": 24, "K": 45}.items():
        sheet.column_dimensions[column].width = width
    for row in range(2, 10002):
        sheet[f"A{row}"].number_format = "@"
        sheet[f"D{row}"].number_format = "@"
        sheet[f"G{row}"].number_format = "yyyy-mm-dd"
        sheet[f"H{row}"].number_format = "yyyy-mm-dd"

    status_validation = DataValidation(type="list", formula1='"' + ','.join(STATES) + '"')
    status_validation.error = "Elegí uno de los cuatro estados de la lista."
    status_validation.errorTitle = "Estado inválido"
    status_validation.showErrorMessage = True
    status_validation.errorStyle = "stop"
    sheet.add_data_validation(status_validation)
    status_validation.add("E2:E10001")

    sheet["J1"] = "Estado permitido"
    sheet["K1"] = "Ejemplo / uso"
    for cell in (sheet["J1"], sheet["K1"]):
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    examples = (
        "Inscripción informada como al día",
        "Inscripción informada como adeudada; único estado contactable",
        "Inscripción nueva informada por la planilla",
        "Inscripción informada como abandonada",
    )
    for row, (state, example) in enumerate(zip(STATES, examples), 2):
        sheet[f"J{row}"] = state
        sheet[f"K{row}"] = example
    sheet["J7"] = "Cómo usar"
    sheet["K7"] = "Completá A:H desde la fila 2. No agregues importes ni fórmulas."
    sheet["K8"] = "Actividad es obligatoria para los cuatro estados."
    sheet["K9"] = "Fechas opcionales: AAAA-MM-DD o DD/MM/AAAA; 'Sin Pagos' equivale a sin vencimiento."
    sheet["K10"] = "DNI y Teléfono como texto. Con ARS: 10 dígitos locales o +54 9 y 10 dígitos."
    sheet["K11"] = "La lista reemplaza la importación CRM vigente del club."
    sheet["K12"] = "Los estados son declarados; no prueban saldos financieros."
    sheet["J14"] = "Versión de plantilla"
    sheet["K14"] = "CRM_CONTACTOS_v1"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


if __name__ == "__main__":
    build()
